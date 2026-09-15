#!/usr/bin/env python3
"""
Turns Hana's hotel spreadsheet (data/hotel-page-data.xlsx: Properties, Room categories, To confirm)
into data/hotels.json, which scripts/build-hotels.mjs reads.

    python3 scripts/import-hotels.py            # rewrites data/hotels.json

The spreadsheet stays the master. Edit the sheet, run this, run the build, commit all three.
Cells that say "TBC" or "Not published" are dropped, so they never reach a page.
The page slug (the address) is the slug the Getaways package data already uses for that hotel,
mapped from the sheet slug below; any hotel not in the map keeps its sheet slug.
Guest-facing copy (intro, good to know) is not in the sheet: it lives in data/hotels-copy.json
and is merged by the build.
"""
import json, re, sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("pip install openpyxl --break-system-packages")

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data/hotel-page-data.xlsx"
OUT = ROOT / "data/hotels.json"

# sheet slug -> page slug (the slug in data/getaways.json), destination slug
MAP = {
    "riu-plaza-panama": ("riu-plaza", "panama"),
    "megapolis-panama": ("megapolis", "panama"),
    "las-americas-golden-tower": ("las-americas", "panama"),
    "marinn-place": ("marinn-place", "panama"),
    "ramada-plaza-panama": ("ramada-plaza", "panama"),
    "hospedium-princess-panama": ("hospedium-princess", "panama"),
    "grand-decameron-panama": ("decameron-panama", "panama"),
    "dreams-playa-bonita": ("dreams-playa-bonita", "panama"),
    "riu-playa-blanca": ("riu-playa-blanca", "panama"),
    "suites-del-bosque": ("suites-del-bosque", "lima"),
    "nobility-hotel": ("nobility", "lima"),
    "nhow-lima": ("nhow", "lima"),
    "souma-lima": ("souma", "lima"),
    "estelar-miraflores": ("estelar-miraflores", "lima"),
    "melia-lima": ("melia-lima", "lima"),
    "hilton-lima-miraflores": ("hilton-lima", "lima"),
}

UNKNOWN = re.compile(r"^\s*(tbc|not published|not shown|on request)\b", re.I)

def clean(v):
    if v is None:
        return None
    s = str(v).strip()
    if not s or UNKNOWN.match(s):
        return None
    return s

def key(field):
    return re.sub(r"[^a-z0-9]+", "_", field.strip().lower()).strip("_")

wb = openpyxl.load_workbook(SRC, data_only=True)
props = list(wb["Properties"].iter_rows(values_only=True))
header = props[0]
hotels = {}
for col, name in enumerate(header):
    if col == 0 or not name:
        continue
    rec = {"name": str(name).strip()}
    for row in props[1:]:
        field = row[0]
        if not field:
            continue
        v = clean(row[col])
        if key(field) in ("latitude", "longitude"):
            try:
                v = float(v) if v is not None else None
            except ValueError:
                v = None
        rec[key(field)] = v
    sheet_slug = rec.get("slug")
    if not sheet_slug:
        continue
    page_slug, dest = MAP.get(sheet_slug, (sheet_slug, None))
    rec["sheetSlug"] = sheet_slug
    rec["slug"] = page_slug
    rec["dest"] = dest
    rec["rooms"] = []
    hotels[sheet_slug] = rec

for row in list(wb["Room categories"].iter_rows(values_only=True))[1:]:
    slug, _hname, rname, size, sleeps, beds, notes = (list(row) + [None] * 7)[:7]
    if not slug or slug not in hotels or not rname:
        continue
    hotels[slug]["rooms"].append({
        "name": str(rname).strip(),
        "size": clean(size),
        "sleeps": clean(sleeps),
        "beds": clean(beds),
        "notes": clean(notes),
    })

# open questions per hotel, kept for the team (never rendered)
open_items = {}
for row in list(wb["To confirm"].iter_rows(values_only=True))[1:]:
    hotel, what, why, who, status = (list(row) + [None] * 5)[:5]
    if not hotel or not what:
        continue
    open_items.setdefault(str(hotel).strip(), []).append({"what": str(what).strip(), "status": clean(status) or "Open"})

out = {
    "source": SRC.name,
    "hotels": sorted(hotels.values(), key=lambda h: (h["dest"] or "zz", h["slug"])),
    "toConfirm": open_items,
}
OUT.write_text(json.dumps(out, ensure_ascii=False, indent=1) + "\n")
n_pages = sum(1 for h in hotels.values() if h["dest"])
print(f"wrote {OUT.relative_to(ROOT)}: {len(hotels)} hotels, {n_pages} with a destination, {sum(len(h['rooms']) for h in hotels.values())} room rows")
