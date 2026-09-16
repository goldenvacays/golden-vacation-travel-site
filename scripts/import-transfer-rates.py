#!/usr/bin/env python3
"""
Turns Hana's transfer workbook (Jamaica_Transfer_Rates_MBJ_KIN_OCJ.xlsx) into data/transfer-rates.json.

Only SELL prices leave this script. Cost, markup and the supplier never reach the JSON, the page or the functions,
and the workbook itself is never committed (it holds cost prices).

  python3 scripts/import-transfer-rates.py path/to/Jamaica_Transfer_Rates_MBJ_KIN_OCJ.xlsx

The model (Sep 16 2026, Hana's "Website Pricing" tab): one price per ZONE, not per hotel. Every hotel in a zone shares the
zone's prices ("deliberate, it keeps the website simple"). A hotel's zone comes from data/transfers.json (hotelZones and
the zones' places lists, plus the region rule for the rest).

Shape of the output:
  zones[zone][airport][trip]   offers for a zone from an airport: "in" (airport to hotel), "out" (hotel to airport, the
                               same one-way price) and "both" (round trip, the sheet's own column). An offer is compact:
                               [rung, seats, bags, price, multi] with multi = 1 when the sheet shows the price in italics,
                               meaning two or more vehicles.
  links["a|b"] / links["a>b"]  hotel-to-hotel prices between two zones ("|" either direction, ">" one direction only),
                               from the Hotel to Hotel tab's SELL block. Round trip is two one-ways (no discount, per the sheet).
  exceptions                   hotel + airport + seat-count combinations the sheet says to quote by hand; the page and the
                               checkout show "priced in the chat" for exactly those.
Passenger tiers become vehicle rungs: 3 car, 5 and 7 minivan, 9, 14 and 16 minibus. The 25 tier is left out (the site
takes parties up to 16; bigger ones go by email). Hotel-to-hotel tiers: 3 car, 4 and 6 minivan, 8, 10, 12 and 16 minibus.
Bags: one per seat is the guide.
"""
import json, re, sys
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else None
if not SRC or not SRC.exists():
    sys.exit("usage: import-transfer-rates.py <xlsx>")

AIRPORT_HEAD = {"Montego Bay (MBJ)": "MBJ", "Ocho Rios (OCJ)": "OCJ", "Kingston (KIN)": "KIN", "Negril Aerodrome (NEG)": None}  # NEG stays off the site
MAX_SEATS = 16

def rung(seats):
    return "car" if seats <= 3 else "minivan" if seats <= 7 else "minibus" if seats <= 16 else "coach"

# sheet zone label (start of the cell) -> site zone keys it prices
ZONE_ROWS = [
    ("Montego Bay - town, Ironshore & Rose Hall", ["mobay", "rose-hall"]),
    ("Montego Bay - not all-inclusive", ["mobay-nonai"]),
    ("Montego Bay – town / Ironshore", ["mobay"]),
    ("Montego Bay – Rose Hall", ["rose-hall"]),
    ("Falmouth / Trelawny", ["falmouth"]),
    ("Runaway Bay", ["runaway-bay"]),
    ("Ocho Rios", ["ocho-rios"]),
    ("Hanover (Lucea / Green Island)", ["lucea"]),
    ("Hanover - villa estates", ["hanover-villas"]),
    ("Negril", ["negril"]),
    ("South Coast", ["south-coast"]),
    ("Treasure Beach / Black River", ["treasure-beach"]),
    ("Port Antonio", ["port-antonio"]),
    ("Kingston - New Kingston / city hotels", ["kingston"]),
    ("Kingston - Blue Mountains", ["blue-mountains"]),
]
def zone_keys(label):
    label = str(label).strip()
    for start, keys in ZONE_ROWS:
        if label.startswith(start):
            return keys
    return None

# hotel-to-hotel route ends -> zone keys (Montego Bay covers the whole Montego Bay side)
LINK_END = {
    "Montego Bay": ["mobay", "rose-hall", "mobay-nonai"], "Ocho Rios": ["ocho-rios"], "Negril": ["negril"], "Runaway Bay": ["runaway-bay"],
    "Falmouth": ["falmouth"], "South Coast (Whitehouse)": ["south-coast"], "Kingston": ["kingston"], "Port Antonio": ["port-antonio"],
    "Treasure Beach (St. E)": ["treasure-beach"],
}

# the "quote these individually" lines: sheet hotel name -> resort-list names
EXC_MATCH = {
    "Excellence Oyster Bay": ["Excellence Oyster Bay"],
    "Ocean Coral Spring & Ocean Eden Bay complex": ["Ocean Coral Spring", "Ocean Eden Bay"],
    "Royalton Negril & Hideaway at Royalton Negril complex": ["Royalton Negril", "Royalton Hideaway Negril"],
    "Sandals Negril": ["Sandals Negril"],
    "Couples Tower Isle": ["Couples Tower Isle"],
    "Riu Ocho Rios": ["RIU Ocho Rios"],
    "Sandals Dunn's River": ["Sandals Dunn's River"],
}

wb = openpyxl.load_workbook(SRC, data_only=True)

# ---------- Website Pricing: zone prices per airport ----------
ws = wb["Website Pricing"]
zones = {}
airport = None
tiers = None  # list of (col index, seats, "one" | "rt")
exceptions_raw = []
in_exceptions = False
for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
    a = row[0].value
    if a is None:
        continue
    a_str = str(a).strip()
    if a_str.startswith("From "):
        airport = None
        for head, code in AIRPORT_HEAD.items():
            if head in a_str:
                airport = code
        tiers = None
        in_exceptions = False
        continue
    if a_str.startswith("Quote these individually"):
        in_exceptions = True
        continue
    if in_exceptions:
        exceptions_raw.append(a_str.rstrip("*").strip())
        continue
    if a_str == "Resort zone":
        tiers = []
        for c in row[1:]:
            if not c.value:
                continue
            m = re.match(r"Up to (\d+) passengers\s*(one-way|round trip)", str(c.value).replace("\n", " "))
            if m:
                tiers.append((c.column - 1, int(m.group(1)), "one" if m.group(2) == "one-way" else "rt"))
        continue
    if airport is None or tiers is None:
        continue
    keys = zone_keys(a_str)
    if not keys:
        continue
    one, rt = [], []
    for col, seats, kind in tiers:
        cell = row[col]
        if cell.value is None or seats > MAX_SEATS:
            continue
        price = int(round(float(cell.value)))
        multi = 1 if (cell.font and cell.font.i) else 0
        offer = [rung(seats), seats, seats, price] + ([1] if multi else [])
        (one if kind == "one" else rt).append(offer)
    if not one:
        continue
    for k in keys:
        zones.setdefault(k, {})[airport] = {"in": one, "out": [list(o) for o in one], "both": rt}

# ---------- Hotel to Hotel: the SELL block ----------
ws = wb["Hotel to Hotel"]
links = {}
sell = False
link_tiers = None
for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
    a = row[0].value
    if a is None:
        continue
    a_str = str(a).strip()
    if a_str.startswith("SELL PRICE"):
        sell = True
        continue
    if not sell:
        continue
    if a_str == "Route":
        link_tiers = []
        for c in row[1:]:
            m = c.value and re.match(r"Up to (\d+) passengers", str(c.value))
            if m:
                link_tiers.append((c.column - 1, int(m.group(1))))
        continue
    if a_str.startswith("Notes") or a_str.startswith("Markup"):
        break
    if link_tiers is None:
        continue
    m = re.match(r"(.+?)\s*(↔|->)\s*(.+)$", a_str)
    if not m:
        continue
    left, arrow, right = m.group(1).strip(), m.group(2), m.group(3).strip()
    if left not in LINK_END or right not in LINK_END:
        print(f"hotel-to-hotel route skipped, unknown end: {a_str}", file=sys.stderr)
        continue
    offers = []
    for col, seats in link_tiers:
        cell = row[col]
        if cell.value is None or seats > MAX_SEATS:
            continue
        price = int(round(float(cell.value)))
        multi = 1 if (cell.font and cell.font.i) else 0
        offers.append([rung(seats), seats, seats, price] + ([1] if multi else []))
    for za in LINK_END[left]:
        for zb in LINK_END[right]:
            if za == zb:
                continue
            key = f"{za}>{zb}" if arrow == "->" else "|".join(sorted([za, zb]))
            links[key] = offers

# ---------- exceptions ----------
exceptions = []
for line in exceptions_raw:
    m = re.match(r"(MBJ|KIN|OCJ)\s*·\s*(.+?)\s*·\s*(?:Minibus|Coach|Minivan|Car)\s+up to\s+(\d+):\s*(.+)$", line)
    if not m:
        print(f"exception line skipped: {line}", file=sys.stderr)
        continue
    ap, zone_label, seats, names = m.group(1), m.group(2), int(m.group(3)), m.group(4)
    keys = zone_keys(zone_label)
    if not keys or seats > MAX_SEATS:
        continue
    hotels = []
    for n in [x.strip() for x in names.split(",")]:
        hotels += EXC_MATCH.get(n, [n])
    for k in keys:
        exceptions.append({"airport": ap, "zone": k, "seats": seats, "hotels": hotels})

out = {
    "source": "Hana's transfer workbook, Website Pricing and Hotel to Hotel tabs, Sep 16 2026. Sell prices only, USD per vehicle, tax included. Regenerate with scripts/import-transfer-rates.py.",
    "airports": ["MBJ", "KIN", "OCJ"],
    "model": "zone",
    "hotels": {},
    "rates": {},
    "zones": zones,
    "links": links,
    "exceptions": exceptions,
}
(ROOT / "data/transfer-rates.json").write_text(json.dumps(out, separators=(",", ":"), ensure_ascii=False) + "\n", encoding="utf-8")
print(f"zones: {len(zones)} priced ({', '.join(sorted(zones))}); hotel-to-hotel pairs: {len(links)}; exceptions: {len(exceptions)}")
for k in sorted(zones):
    print(f"  {k:16} " + "  ".join(f"{ap}: {len(zones[k][ap]['in'])} sizes, car {zones[k][ap]['in'][0][3]}" for ap in ("MBJ", "KIN", "OCJ") if ap in zones[k]))
